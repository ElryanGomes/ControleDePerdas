"""
app.py — Controle de Perdas (backend Flask + SQLAlchemy)

Rodar localmente:
    pip install -r requirements.txt
    python app.py
    -> abre em http://127.0.0.1:5000

Deploy no PythonAnywhere: veja README.md
"""
import os
import re
import sqlite3
from datetime import datetime
from io import BytesIO

from flask import Flask, request, jsonify, render_template, send_file
from sqlalchemy import func, inspect, text
from sqlalchemy.exc import IntegrityError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import db, Fornecedor, Produto, Perda, Mes, mes_ref_de, mes_seguinte

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)

# Por padrão usa SQLite (arquivo local). Para usar MySQL no PythonAnywhere,
# defina a variável de ambiente DATABASE_URL, por exemplo:
#   mysql://usuario:senha@usuario.mysql.pythonanywhere-services.com/usuario$controle_perdas
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get(
    'DATABASE_URL', 'sqlite:///' + os.path.join(BASE_DIR, 'controle_perdas.db')
)
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

REGEX_MES = re.compile(r'^\d{4}-(0[1-9]|1[0-2])$')


# ================= migração do banco (sistema mensal) =================
def _caminho_sqlite():
    uri = app.config['SQLALCHEMY_DATABASE_URI']
    if not uri.startswith('sqlite:///'):
        return None
    caminho = uri[len('sqlite:///'):]
    if not os.path.isabs(caminho):
        caminho = os.path.join(BASE_DIR, caminho)
    return caminho


def fazer_backup_sqlite():
    """Copia o arquivo .db antes de alterar a estrutura (só SQLite;
    no MySQL, faça um dump manual antes). Retorna o caminho do backup."""
    caminho = _caminho_sqlite()
    if not caminho or not os.path.exists(caminho):
        return None
    carimbo = datetime.now().strftime('%Y%m%d_%H%M%S')
    destino = os.path.splitext(caminho)[0] + f'_{carimbo}.bak'
    origem = sqlite3.connect(caminho)
    copia = sqlite3.connect(destino)
    try:
        origem.backup(copia)
    finally:
        copia.close()
        origem.close()
    print(f'[migração] backup criado: {destino}')
    return destino


def migrar_banco():
    """Adiciona mes_ref às perdas, cria a tabela de meses e define o mês aberto.
    É idempotente (pode rodar a cada inicialização) e NUNCA apaga nem altera
    registros existentes — só preenche o campo novo."""
    backup = None
    try:
        # 1) coluna nova em banco antigo (create_all não altera tabelas existentes)
        insp = inspect(db.engine)
        if 'perdas' in insp.get_table_names():
            colunas = [c['name'] for c in insp.get_columns('perdas')]
            if 'mes_ref' not in colunas:
                backup = fazer_backup_sqlite()
                with db.engine.begin() as conn:
                    conn.execute(text('ALTER TABLE perdas ADD COLUMN mes_ref VARCHAR(7)'))

        # 2) cria o que falta (tabela meses; tudo, se o banco for novo)
        db.create_all()

        insp = inspect(db.engine)
        if 'ix_perdas_mes_ref' not in [i['name'] for i in insp.get_indexes('perdas')]:
            with db.engine.begin() as conn:
                conn.execute(text('CREATE INDEX ix_perdas_mes_ref ON perdas (mes_ref)'))

        # 3) preenche mes_ref das perdas antigas pelo criado_em (horário de Brasília)
        nulos = Perda.query.filter(Perda.mes_ref.is_(None)).all()
        if nulos and backup is None:
            backup = fazer_backup_sqlite()
        ref_agora = mes_ref_de(datetime.utcnow())
        for perda in nulos:
            perda.mes_ref = mes_ref_de(perda.criado_em) if perda.criado_em else ref_agora
        db.session.commit()

        # 4) tabela de meses: o mais recente com registro fica ABERTO, os anteriores FECHADOS
        existentes = {m.ref: m for m in Mes.query.all()}
        abertos = [m for m in existentes.values() if m.status == 'ABERTO']
        if len(abertos) > 1:  # defesa: só um pode ficar aberto
            abertos.sort(key=lambda m: m.ref)
            for m in abertos[:-1]:
                m.status = 'FECHADO'
            abertos = abertos[-1:]

        refs_perdas = {r[0] for r in db.session.query(Perda.mes_ref).distinct() if r[0]}
        todos = refs_perdas | set(existentes)
        ref_aberto = abertos[0].ref if abertos else (max(todos) if todos else ref_agora)
        agora = datetime.utcnow()

        for ref in sorted(todos | {ref_aberto}):
            deve_abrir = (ref == ref_aberto) and not abertos
            if ref not in existentes:
                db.session.add(Mes(ref=ref, status='ABERTO' if deve_abrir else 'FECHADO',
                                   aberto_em=agora if deve_abrir else None))
            elif deve_abrir:
                existentes[ref].status = 'ABERTO'
                existentes[ref].fechado_em = None
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        aviso = f' Backup do banco: {backup}.' if backup else ''
        raise RuntimeError(f'Falha na migração do banco: {e}.{aviso} Nenhum registro foi apagado.') from e


with app.app_context():
    migrar_banco()


def mes_aberto():
    return Mes.query.filter_by(status='ABERTO').first()


def garantir_mes_aberto():
    """Devolve o mês aberto; se por algum motivo não houver, abre o mês corrente."""
    aberto = mes_aberto()
    if aberto:
        return aberto
    ref = mes_ref_de(datetime.utcnow())
    existente = Mes.query.filter_by(ref=ref).first()
    if existente:
        existente.status = 'ABERTO'
        existente.fechado_em = None
        aberto = existente
    else:
        aberto = Mes(ref=ref, status='ABERTO', aberto_em=datetime.utcnow())
        db.session.add(aberto)
    db.session.flush()
    return aberto


def bloquear_se_fechado(perda):
    """Retorna uma resposta 403 se a perda pertence a um mês fechado, senão None."""
    mes = Mes.query.filter_by(ref=perda.mes_ref).first() if perda.mes_ref else None
    if mes and mes.status == 'FECHADO':
        return jsonify({'erro': f'O mês {perda.mes_ref} está fechado — os registros dele são somente leitura.'}), 403
    return None


def get_or_create_fornecedor(nome):
    nome = (nome or '').strip()
    fornecedor = Fornecedor.query.filter_by(nome=nome).first()
    if not fornecedor:
        fornecedor = Fornecedor(nome=nome)
        db.session.add(fornecedor)
        db.session.flush()  # garante fornecedor.id antes do commit
    return fornecedor


# ================= páginas =================
@app.route('/')
def index():
    return render_template('index.html')


# ================= fornecedores =================
@app.route('/api/fornecedores', methods=['GET'])
def listar_fornecedores():
    fornecedores = Fornecedor.query.order_by(Fornecedor.nome).all()
    return jsonify([f.nome for f in fornecedores])


@app.route('/api/fornecedores', methods=['POST'])
def criar_fornecedor():
    data = request.get_json(force=True) or {}
    nome = (data.get('nome') or '').strip()
    if not nome:
        return jsonify({'erro': 'Nome do fornecedor é obrigatório.'}), 400
    get_or_create_fornecedor(nome)
    db.session.commit()
    return jsonify({'ok': True})


# Ao renomear um fornecedor, as perdas já lançadas (que guardam o nome como texto)
# também passam a usar o novo nome, para Relatório e Rankings não separarem o mesmo
# fornecedor em dois. True = inclui também os meses FECHADOS. Se quiser que os meses
# fechados fiquem exatamente como foram fechados, troque para False.
RENOMEAR_EM_MESES_FECHADOS = True


@app.route('/api/fornecedores/<path:nome>', methods=['PUT'])
def atualizar_fornecedor(nome):
    """Renomeia um fornecedor. Os produtos acompanham (são o mesmo registro) e as
    perdas lançadas com o nome antigo recebem o novo."""
    fornecedor = Fornecedor.query.filter_by(nome=(nome or '').strip()).first()
    if not fornecedor:
        return jsonify({'erro': 'Fornecedor não encontrado.'}), 404

    data = request.get_json(force=True) or {}
    novo_nome = (data.get('nome') or '').strip()
    if not novo_nome:
        return jsonify({'erro': 'Nome do fornecedor é obrigatório.'}), 400
    if len(novo_nome) > 200:
        return jsonify({'erro': 'Nome do fornecedor muito longo (máximo 200 caracteres).'}), 400

    # não deixa criar dois fornecedores iguais (ignorando maiúsculas/minúsculas);
    # mudar só a caixa do próprio fornecedor (ex: Nestle -> NESTLE) é permitido
    duplicado = Fornecedor.query.filter(
        func.lower(Fornecedor.nome) == novo_nome.lower(),
        Fornecedor.id != fornecedor.id,
    ).first()
    if duplicado:
        return jsonify({'erro': f'Já existe um fornecedor chamado "{duplicado.nome}".'}), 409

    nome_antigo = fornecedor.nome
    fornecedor.nome = novo_nome

    perdas = Perda.query.filter(Perda.fornecedor == nome_antigo)
    if not RENOMEAR_EM_MESES_FECHADOS:
        fechados = [m.ref for m in Mes.query.filter_by(status='FECHADO').all()]
        if fechados:
            perdas = perdas.filter(~Perda.mes_ref.in_(fechados))
    atualizadas = perdas.update({Perda.fornecedor: novo_nome}, synchronize_session=False)

    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({'erro': 'Já existe um fornecedor com esse nome.'}), 409
    return jsonify({'ok': True, 'nome': novo_nome, 'perdasAtualizadas': atualizadas})


@app.route('/api/fornecedores/<path:nome>', methods=['DELETE'])
def excluir_fornecedor(nome):
    """Apaga o fornecedor só se ele não tiver produtos. O histórico de perdas
    continua intacto (as perdas guardam o nome do fornecedor como texto)."""
    fornecedor = Fornecedor.query.filter_by(nome=(nome or '').strip()).first()
    if not fornecedor:
        return jsonify({'erro': 'Fornecedor não encontrado.'}), 404

    qtd = Produto.query.filter_by(fornecedor_id=fornecedor.id).count()
    if qtd:
        return jsonify({'erro': f'"{fornecedor.nome}" ainda tem {qtd} produto(s) cadastrado(s). '
                                'Apague ou mude esses produtos antes.'}), 409

    db.session.delete(fornecedor)
    db.session.commit()
    return jsonify({'ok': True})


# ================= catálogo (fornecedores + produtos agrupados) =================
@app.route('/api/catalogo', methods=['GET'])
def obter_catalogo():
    fornecedores = Fornecedor.query.order_by(Fornecedor.nome).all()
    catalogo = {f.nome: [p.to_dict() for p in f.produtos] for f in fornecedores}
    return jsonify(catalogo)


# ================= produtos =================
@app.route('/api/produtos', methods=['POST'])
def upsert_produto():
    """Cria o produto, ou atualiza se já existir (por código de barras ou código do produto,
    dentro do mesmo fornecedor)."""
    data = request.get_json(force=True) or {}
    fornecedor_nome = (data.get('fornecedor') or '').strip()
    descricao = (data.get('descricao') or '').strip()
    if not fornecedor_nome:
        return jsonify({'erro': 'Fornecedor é obrigatório.'}), 400
    if not descricao:
        return jsonify({'erro': 'Descrição é obrigatória.'}), 400

    fornecedor = get_or_create_fornecedor(fornecedor_nome)
    codigo_barra = (data.get('codigoBarra') or '').strip()
    codigo_produto = (data.get('codigoProduto') or '').strip()
    valor_unit = float(data.get('valorUnit') or 0)

    produto = None
    if codigo_barra:
        produto = Produto.query.filter_by(fornecedor_id=fornecedor.id, codigo_barra=codigo_barra).first()
    if not produto and codigo_produto:
        produto = Produto.query.filter_by(fornecedor_id=fornecedor.id, codigo_produto=codigo_produto).first()

    if produto:
        produto.codigo_barra = codigo_barra
        produto.codigo_produto = codigo_produto
        produto.descricao = descricao
        produto.valor_unit = valor_unit
    else:
        produto = Produto(
            fornecedor_id=fornecedor.id, codigo_barra=codigo_barra,
            codigo_produto=codigo_produto, descricao=descricao, valor_unit=valor_unit,
        )
        db.session.add(produto)

    db.session.commit()
    return jsonify(produto.to_dict())


@app.route('/api/produtos/<int:produto_id>', methods=['PUT'])
def atualizar_produto(produto_id):
    produto = Produto.query.get_or_404(produto_id)
    data = request.get_json(force=True) or {}

    fornecedor_nome = (data.get('fornecedor') or '').strip()
    if fornecedor_nome:
        fornecedor = get_or_create_fornecedor(fornecedor_nome)
        produto.fornecedor_id = fornecedor.id

    descricao = (data.get('descricao') or '').strip()
    if not descricao:
        return jsonify({'erro': 'Descrição é obrigatória.'}), 400

    produto.codigo_barra = (data.get('codigoBarra') or '').strip()
    produto.codigo_produto = (data.get('codigoProduto') or '').strip()
    produto.descricao = descricao
    produto.valor_unit = float(data.get('valorUnit') or 0)

    db.session.commit()
    return jsonify(produto.to_dict())


@app.route('/api/produtos/<int:produto_id>', methods=['DELETE'])
def excluir_produto(produto_id):
    produto = Produto.query.get_or_404(produto_id)
    db.session.delete(produto)
    db.session.commit()
    return jsonify({'ok': True})


# ================= perdas =================
@app.route('/api/perdas', methods=['GET'])
def listar_perdas():
    query = Perda.query
    fornecedor = request.args.get('fornecedor')
    tipo = request.args.get('tipo')
    mes = (request.args.get('mes') or '').strip()
    if mes.lower() == 'todos':
        pass  # todos os meses
    elif mes:
        if not REGEX_MES.match(mes):
            return jsonify({'erro': 'Mês inválido. Use AAAA-MM ou "todos".'}), 400
        query = query.filter_by(mes_ref=mes)
    else:
        aberto = mes_aberto()  # sem parâmetro: só o mês aberto
        query = query.filter_by(mes_ref=aberto.ref if aberto else '0000-00')
    if fornecedor:
        query = query.filter_by(fornecedor=fornecedor)
    if tipo:
        query = query.filter_by(tipo=tipo)
    perdas = query.order_by(Perda.criado_em.desc()).all()
    return jsonify([p.to_dict() for p in perdas])


@app.route('/api/perdas', methods=['POST'])
def criar_perda():
    data = request.get_json(force=True) or {}
    fornecedor = (data.get('fornecedor') or '').strip()
    descricao = (data.get('descricao') or '').strip()
    tipo = (data.get('tipo') or '').strip().upper()
    validade = (data.get('validade') or '').strip()

    try:
        quantidade = int(data.get('quantidade') or 0)
        valor_unit = float(data.get('valorUnit') or 0)
    except (TypeError, ValueError):
        return jsonify({'erro': 'Quantidade ou valor inválido.'}), 400

    if not fornecedor or not descricao or tipo not in ('VENCIDO', 'AVARIADO') or quantidade < 1 or not validade:
        return jsonify({'erro': 'Preencha todos os campos obrigatórios.'}), 400

    get_or_create_fornecedor(fornecedor)

    aberto = garantir_mes_aberto()
    perda = Perda(
        mes_ref=aberto.ref,
        fornecedor=fornecedor,
        descricao=descricao,
        codigo=(data.get('codigo') or '').strip(),
        tipo=tipo,
        quantidade=quantidade,
        validade=validade,
        valor_unit=valor_unit,
        valor_total=round(quantidade * valor_unit, 2),
    )
    db.session.add(perda)
    db.session.commit()
    return jsonify(perda.to_dict())


@app.route('/api/perdas/<int:perda_id>', methods=['PUT'])
def atualizar_perda(perda_id):
    perda = Perda.query.get_or_404(perda_id)
    bloqueio = bloquear_se_fechado(perda)
    if bloqueio:
        return bloqueio
    data = request.get_json(force=True) or {}

    fornecedor = (data.get('fornecedor') or '').strip()
    descricao = (data.get('descricao') or '').strip()
    tipo = (data.get('tipo') or '').strip().upper()
    validade = (data.get('validade') or '').strip()

    try:
        quantidade = int(data.get('quantidade') or 0)
        valor_unit = float(data.get('valorUnit') or 0)
    except (TypeError, ValueError):
        return jsonify({'erro': 'Quantidade ou valor inválido.'}), 400

    if not fornecedor or not descricao or tipo not in ('VENCIDO', 'AVARIADO') or quantidade < 1 or not validade:
        return jsonify({'erro': 'Preencha todos os campos obrigatórios.'}), 400

    get_or_create_fornecedor(fornecedor)

    perda.fornecedor = fornecedor
    perda.descricao = descricao
    perda.codigo = (data.get('codigo') or '').strip()
    perda.tipo = tipo
    perda.quantidade = quantidade
    perda.validade = validade
    perda.valor_unit = valor_unit
    perda.valor_total = round(quantidade * valor_unit, 2)

    db.session.commit()
    return jsonify(perda.to_dict())


@app.route('/api/export/xlsx', methods=['POST'])
def exportar_xlsx():
    """Gera um relatório Excel (.xlsx) de perdas, com uma aba por fornecedor.

    Recebe no corpo JSON os registros já filtrados pelo front-end
    (mesma lista exibida na tela de Relatório), no formato:
        { "perdas": [ {fornecedor, descricao, tipo, quantidade,
                        codigo, validade, valorUnit, valorTotal}, ... ] }
    """
    data = request.get_json(force=True) or {}
    perdas = data.get('perdas') or []
    if not perdas:
        return jsonify({'erro': 'Nenhum registro para exportar.'}), 400

    # agrupa por fornecedor, preservando a ordem de primeira aparição
    grupos = {}
    ordem = []
    for p in perdas:
        forn = (p.get('fornecedor') or 'SEM FORNECEDOR').strip()
        if forn not in grupos:
            grupos[forn] = []
            ordem.append(forn)
        grupos[forn].append(p)

    wb = Workbook()
    wb.remove(wb.active)  # remove a aba padrão em branco

    # ---------- estilos ----------
    AZUL_ESCURO = 'FF1F3864'
    AZUL_TITULO = 'FF2F75B5'
    AZUL_CLARO = 'FFDCE6F1'
    BRANCO = 'FFFFFFFF'
    CINZA_BORDA = 'FFB7B7B7'

    fonte_header_principal = Font(name='Calibri', size=14, bold=True, color=AZUL_ESCURO)
    fonte_titulo_coluna = Font(name='Calibri', size=11, bold=True, color=BRANCO)
    fonte_rodape = Font(name='Calibri', size=11, bold=True)

    fill_titulo = PatternFill('solid', fgColor=AZUL_TITULO)
    fill_soma = PatternFill('solid', fgColor=AZUL_CLARO)

    borda_fina = Border(
        left=Side(style='thin', color=CINZA_BORDA),
        right=Side(style='thin', color=CINZA_BORDA),
        top=Side(style='thin', color=CINZA_BORDA),
        bottom=Side(style='thin', color=CINZA_BORDA),
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    COLUNAS = ['ITEM', 'DESCRIÇÃO', 'STATUS', 'QUANT.', 'CÓDIGO', 'VALIDADE', 'VALOR UNIT.', 'VALOR TOTAL']
    LARGURAS = [8, 45, 14, 10, 14, 14, 14, 16]
    FORMATO_MOEDA = 'R$ #,##0.00'

    nomes_usados = set()

    for forn in ordem:
        itens = grupos[forn]

        # nome da aba: versão curta/título do fornecedor (ex: "M Dias Branco")
        nome_aba = forn.strip().title()
        for ch in ('\\', '/', '?', '*', '[', ']', ':'):
            nome_aba = nome_aba.replace(ch, '')
        nome_aba = (nome_aba[:31] or 'Fornecedor')
        base_nome, sufixo_n = nome_aba, 2
        while nome_aba in nomes_usados:
            sufixo = f' ({sufixo_n})'
            nome_aba = base_nome[:31 - len(sufixo)] + sufixo
            sufixo_n += 1
        nomes_usados.add(nome_aba)

        ws = wb.create_sheet(title=nome_aba)

        for idx, largura in enumerate(LARGURAS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = largura

        # ---------- 1) cabeçalho principal (linha 1) ----------
        ws.merge_cells('A1:H1')
        titulo_cell = ws['A1']
        titulo_cell.value = f'RELATÓRIO DE PERDAS - {forn.strip().upper()}'
        titulo_cell.font = fonte_header_principal
        titulo_cell.alignment = align_center
        ws.row_dimensions[1].height = 26

        # ---------- 2) títulos das colunas (linha 3) ----------
        header_row = 3
        for col_idx, titulo in enumerate(COLUNAS, start=1):
            c = ws.cell(row=header_row, column=col_idx, value=titulo)
            c.font = fonte_titulo_coluna
            c.fill = fill_titulo
            c.alignment = align_center
            c.border = borda_fina

        # ---------- 3) dados (a partir da linha 4) ----------
        primeira_linha = header_row + 1
        linha = primeira_linha
        for item_num, p in enumerate(itens, start=1):
            quantidade = p.get('quantidade') or 0
            valor_unit = p.get('valorUnit') or 0

            valores = [
                item_num,
                p.get('descricao') or '',
                p.get('tipo') or '',
                quantidade,
                p.get('codigo') or '',
                p.get('validade') or '',
                valor_unit,
                None,  # VALOR TOTAL vira fórmula abaixo
            ]
            for col_idx, valor in enumerate(valores, start=1):
                c = ws.cell(row=linha, column=col_idx, value=valor)
                c.border = borda_fina
                c.alignment = align_left if col_idx == 2 else align_center

            ws.cell(row=linha, column=7).number_format = FORMATO_MOEDA

            formula_cell = ws.cell(row=linha, column=8, value=f'=D{linha}*G{linha}')
            formula_cell.number_format = FORMATO_MOEDA
            formula_cell.border = borda_fina
            formula_cell.alignment = align_center

            linha += 1

        ultima_linha_dados = linha - 1

        # ---------- 4) rodapé de totalização ----------
        linha_rodape = linha
        ws.merge_cells(f'A{linha_rodape}:G{linha_rodape}')
        rotulo = ws.cell(row=linha_rodape, column=1, value='TOTAL DE PERDAS DESTE FORNECEDOR:')
        rotulo.font = fonte_rodape
        rotulo.alignment = align_right
        for col_idx in range(1, 8):
            ws.cell(row=linha_rodape, column=col_idx).border = borda_fina

        soma_cell = ws.cell(
            row=linha_rodape, column=8,
            value=f'=SUM(H{primeira_linha}:H{ultima_linha_dados})'
        )
        soma_cell.font = fonte_rodape
        soma_cell.number_format = FORMATO_MOEDA
        soma_cell.fill = fill_soma
        soma_cell.alignment = align_center
        soma_cell.border = borda_fina

        ws.freeze_panes = f'A{primeira_linha}'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    mes_arq = (data.get('mes') or '').strip()
    sufixo = mes_arq if REGEX_MES.match(mes_arq) else datetime.utcnow().strftime('%Y-%m-%d')
    filename = f"relatorio_perdas_{sufixo}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/api/perdas/<int:perda_id>', methods=['DELETE'])
def excluir_perda(perda_id):
    perda = Perda.query.get_or_404(perda_id)
    bloqueio = bloquear_se_fechado(perda)
    if bloqueio:
        return bloqueio
    db.session.delete(perda)
    db.session.commit()
    return jsonify({'ok': True})


# ================= meses =================
@app.route('/api/meses', methods=['GET'])
def listar_meses():
    """Meses com status e totais (valor e quantidade) de vencido/avariado."""
    linhas = db.session.query(
        Perda.mes_ref, Perda.tipo,
        func.count(Perda.id),
        func.coalesce(func.sum(Perda.quantidade), 0),
        func.coalesce(func.sum(Perda.valor_total), 0),
    ).group_by(Perda.mes_ref, Perda.tipo).all()

    totais = {}
    for ref, tipo, registros, qtd, valor in linhas:
        t = totais.setdefault(ref, {'registros': 0,
                                    'VENCIDO': {'valor': 0.0, 'quantidade': 0},
                                    'AVARIADO': {'valor': 0.0, 'quantidade': 0}})
        t['registros'] += registros
        if tipo in ('VENCIDO', 'AVARIADO'):
            t[tipo]['valor'] = round(float(valor), 2)
            t[tipo]['quantidade'] = int(qtd)

    meses = []
    for m in Mes.query.order_by(Mes.ref.desc()).all():
        t = totais.get(m.ref) or {'registros': 0,
                                  'VENCIDO': {'valor': 0.0, 'quantidade': 0},
                                  'AVARIADO': {'valor': 0.0, 'quantidade': 0}}
        item = m.to_dict()
        item['registros'] = t['registros']
        item['vencido'] = t['VENCIDO']
        item['avariado'] = t['AVARIADO']
        meses.append(item)

    aberto = mes_aberto()
    return jsonify({
        'aberto': aberto.ref if aberto else None,
        'calendario': mes_ref_de(datetime.utcnow()),
        'meses': meses,
    })


@app.route('/api/meses/fechar', methods=['POST'])
def fechar_mes():
    """Fecha o mês aberto (vira somente leitura) e abre o seguinte.
    Nada é apagado. Recebe {"mes": "AAAA-MM"} para evitar fechar duas vezes."""
    data = request.get_json(force=True) or {}
    ref = (data.get('mes') or '').strip()
    aberto = mes_aberto()
    if not aberto or aberto.ref != ref:
        return jsonify({'erro': 'Esse mês já foi fechado ou não é o mês aberto. Atualize a página.'}), 409

    agora = datetime.utcnow()
    aberto.status = 'FECHADO'
    aberto.fechado_em = agora

    proximo_ref = mes_seguinte(ref)
    proximo = Mes.query.filter_by(ref=proximo_ref).first()
    if proximo:
        proximo.status = 'ABERTO'
        proximo.aberto_em = proximo.aberto_em or agora
        proximo.fechado_em = None
    else:
        db.session.add(Mes(ref=proximo_ref, status='ABERTO', aberto_em=agora))

    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({'erro': 'Esse mês já foi fechado. Atualize a página.'}), 409
    return jsonify({'fechado': ref, 'aberto': proximo_ref})


@app.route('/api/fornecedores/uso', methods=['GET'])
def uso_fornecedores():
    """Quantidade de perdas por fornecedor (todos os meses) — ordena os atalhos do Registrar."""
    linhas = db.session.query(Perda.fornecedor, func.count(Perda.id)).group_by(Perda.fornecedor).all()
    return jsonify({nome: n for nome, n in linhas})


if __name__ == '__main__':
    app.run(debug=True)